from datetime import date
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from constants import (
    EXCEL_PATH,
    SOURCE_URL,
    VARIANT_NAME,
    CLASS_NUMBER,
    YEAR,
    BASE_DIR,
)
from parser.gemini import (
    convert_condition_to_latex,
    QuotaWaitTooLong,
    latex_cache,
)
from parser.latex_cleaner import clean_latex


def load_answers():
    answers_path = BASE_DIR / 'answers.json'

    if not answers_path.exists():
        return {}

    try:
        with answers_path.open('r', encoding='utf-8') as file:
            return json.load(file)
    except (json.JSONDecodeError, OSError) as error:
        print(f'Не удалось прочитать ответы: {error}')
        return {}


def save_workbook(workbook):
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(EXCEL_PATH)


def create_excel(tasks):
    workbook = Workbook()
    tasks_sheet = workbook.active
    tasks_sheet.title = 'Tasks'

    tasks_sheet.append([
        'task_num',
        'condition',
        'image_name',
        'solution',
        'answer',
    ])

    answers = load_answers()
    stop_gemini = False

    for task in tasks:
        task_num = task['task_num']
        condition = clean_latex(task.get('condition', ''), task_num)
        cache_key = f'{task_num}:{condition}'

        if cache_key in latex_cache:
            condition_latex = latex_cache[cache_key]
            print(f'Задание {task_num}: использую кэш.')

        elif stop_gemini:
            condition_latex = condition

        else:
            try:
                condition_latex = convert_condition_to_latex(
                    condition,
                    task_num,
                )
            except QuotaWaitTooLong as error:
                print(f'\n{error}')
                print('Останавливаю запросы Gemini.')
                condition_latex = condition
                stop_gemini = True
            except Exception as error:
                print(f'Ошибка Gemini для задания {task_num}: {error}')
                condition_latex = condition

        condition_latex = clean_latex(condition_latex, task_num)

        answer = task.get('answer', '')
        if not answer:
            answer = answers.get(
                str(task_num),
                answers.get(task_num, ''),
            )

        tasks_sheet.append([
            task_num,
            condition_latex,
            task.get('image_name', ''),
            task.get('solution', ''),
            answer,
        ])

        # Сохраняем после каждого задания, чтобы не потерять прогресс.
        save_workbook(workbook)

    about_sheet = workbook.create_sheet('about')

    metadata = [
        ['Метаданные', 'Значение'],
        ['Класс', CLASS_NUMBER],
        ['Год', YEAR],
        ['Дата создания', date.today().isoformat()],
        ['Название варианта', VARIANT_NAME],
        ['Ссылка на источник', SOURCE_URL],
    ]

    for row in metadata:
        about_sheet.append(row)

    header_font = Font(bold=True)
    header_alignment = Alignment(
        horizontal='center',
        vertical='center',
        wrap_text=True,
    )
    body_alignment = Alignment(
        horizontal='left',
        vertical='top',
        wrap_text=True,
    )

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = body_alignment

        for cell in sheet[1]:
            cell.font = header_font
            cell.alignment = header_alignment

        if sheet.title == 'Tasks':
            widths = {
                'A': 12,
                'B': 75,
                'C': 28,
                'D': 35,
                'E': 25,
            }
            sheet.row_dimensions[1].height = 30
        else:
            widths = {'A': 24, 'B': 55}

        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    save_workbook(workbook)

    print('Excel сохранён:', EXCEL_PATH)
    print('Листы файла:', workbook.sheetnames)


if __name__ == '__main__':
    create_excel([])